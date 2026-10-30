import os
import tempfile
from urllib.parse import urlparse

from faster_whisper import WhisperModel
from langchain_core.runnables import RunnableBranch, RunnableLambda
from langchain_community.document_loaders import (
    PyMuPDFLoader,  # Fast, accurate PDF loader with metadata
    Docx2txtLoader,  # Lightweight DOCX loader
    TextLoader,  # Standard TXT loader
    UnstructuredPowerPointLoader,  # Handles both .ppt and .pptx
    WebBaseLoader,  # Standard web page scraper
    WikipediaLoader,  # Direct Wikipedia API loader
)
from langchain_community.document_loaders.generic import GenericLoader
from langchain_community.document_loaders.parsers import LanguageParser

# Initialize Faster-Whisper once outside the branch (adjust device/compute_type as needed)
whisper_model = WhisperModel("base", device="cpu", compute_type="int8")

SOURCE_CODE_EXTENSIONS = {
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".cpp",
    ".c",
    ".cs",
    ".java",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".html",
    ".css",
}

# --- Loader Helper Functions ---


def load_wav_with_whisper(file_path: str):
    """Transcribes WAV using faster-whisper, then loads via TextLoader."""
    segments, info = whisper_model.transcribe(file_path, beam_size=5)
    transcript = " ".join(segment.text.strip() for segment in segments)

    # Write transcript to a temp .txt file and load with TextLoader
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".txt", delete=False, encoding="utf-8"
    ) as tmp:
        tmp.write(transcript)
        tmp_path = tmp.name

    try:
        docs = TextLoader(tmp_path, encoding="utf-8").load()
        # Preserve original audio filename in metadata for the Vector DB
        for doc in docs:
            doc.metadata["source"] = file_path
            doc.metadata["language"] = info.language
        return docs
    finally:
        os.remove(tmp_path)


def load_source_code(file_path: str):
    """Loads source code using LanguageParser for syntax-aware parsing."""
    loader = GenericLoader.from_filesystem(
        os.path.dirname(file_path) or ".",
        glob=os.path.basename(file_path),
        suffixes=list(SOURCE_CODE_EXTENSIONS),
        parser=LanguageParser(),
    )
    return loader.load()


def load_wikipedia(input_str: str):
    """Extracts the title if a Wikipedia URL is passed, or uses the raw query string."""
    if "wikipedia.org/wiki/" in input_str:
        query = input_str.split("wikipedia.org/wiki/")[-1].replace("_", " ")
    else:
        query = input_str.removeprefix("wiki:").strip()
    return WikipediaLoader(query=query, load_max_docs=2).load()


def load_powerpoint(file_path: str):
    """Loads PPT/PPTX using python-pptx with fallback to UnstructuredPowerPointLoader."""
    try:
        from pptx import Presentation
        from langchain_core.documents import Document

        prs = Presentation(file_path)
        slides_text = []
        for idx, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        t = p.text.strip()
                        if t:
                            texts.append(t)
            if texts:
                slides_text.append(f"--- Slide {idx} ---\n" + "\n".join(texts))

        content = "\n\n".join(slides_text)
        return [Document(page_content=content, metadata={"source": file_path})]
    except Exception:
        return UnstructuredPowerPointLoader(file_path).load()


# --- Routing Conditions ---


def is_wikipedia(x: str) -> bool:
    return "wikipedia.org/wiki/" in x.lower() or x.lower().startswith("wiki:")


def is_web_url(x: str) -> bool:
    parsed = urlparse(x)
    return parsed.scheme in ("http", "https") and not is_wikipedia(x)


def has_ext(x: str, extensions: set[str]) -> bool:
    _, ext = os.path.splitext(x.lower())
    return ext in extensions


# --- RunnableBranch Definition ---

document_loader_branch = RunnableBranch(
    # 1. Wikipedia pages (URL or "wiki:Topic" prefix)
    (
        lambda x: is_wikipedia(x),
        RunnableLambda(load_wikipedia),
    ),
    # 2. General Web URLs
    (
        lambda x: is_web_url(x),
        RunnableLambda(lambda x: WebBaseLoader(x).load()),
    ),
    # 3. PDF Documents
    (
        lambda x: has_ext(x, {".pdf"}),
        RunnableLambda(lambda x: PyMuPDFLoader(x).load()),
    ),
    # 4. DOCX Documents
    (
        lambda x: has_ext(x, {".docx"}),
        RunnableLambda(lambda x: Docx2txtLoader(x).load()),
    ),
    # 5. TXT Files
    (
        lambda x: has_ext(x, {".txt"}),
        RunnableLambda(lambda x: TextLoader(x, encoding="utf-8").load()),
    ),
    # 6. PowerPoint Presentations (PPT / PPTX)
    (
        lambda x: has_ext(x, {".ppt", ".pptx"}),
        RunnableLambda(load_powerpoint),
    ),
    # 7. WAV Audio Files (Faster-Whisper -> TextLoader)
    (
        lambda x: has_ext(x, {".wav"}),
        RunnableLambda(load_wav_with_whisper),
    ),
    # 8. Source-Code Files
    (
        lambda x: has_ext(x, SOURCE_CODE_EXTENSIONS),
        RunnableLambda(load_source_code),
    ),
    # Default Fallback
    RunnableLambda(lambda x: ValueError(f"Unsupported input or file format: {x}")),
)
